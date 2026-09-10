"""Metric registry (ORDER) plus the writers for the per-run long CSV and runs/results.xlsx.
"""
from __future__ import annotations
import csv
import math
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# headline block, in this order
KEY_METRICS = ["acc_prior", "nll_raw", "ece_raw", "T_star", "nll_cal", "ece_cal"]

# bookkeeping keys that are not written to either output (they stay in the per-run JSON)
EXCLUDE = {"status", "figure", "figure_uq", "figure_tree"}

# (key or "prefix*", tier, group, applies, description)   tier 0 = Config section
ORDER = [
    # ---------------- Config ----------------
    ("run_id",      0, "run", "all", "Run identifier (row in the Runs sheet)."),
    ("timestamp",   0, "run", "all",
     "When this record was produced, yymmdd_hh_mm_ss (sorts chronologically as text). "
     "Also names the run folder: runs/<run_id>_<timestamp>/."),
    ("notes",       0, "run", "all", "The free-text 'notes' cell of the run's row in the Runs sheet."),
    ("task",        0, "run", "all",
     "Task: blob | routing | fashion | cifar10_resnet."),
    ("family",      0, "run", "all", "Recognition family: tree | mf | tree_max_span."),
    ("estimator",   0, "run", "all",
     "Reconstruction-gradient estimator that ACTUALLY ran: sfe-loo (score function + leave-one-out) | gumbel (concrete relaxation; ancestral through the tree). CRITICAL for reading a tree-vs-mf comparison: the two families must share this value, otherwise the difference is the estimator, not the recognition family."),
    ("seed",        0, "run", "all", "Random seed for this training run."),
    ("wall_time_s", 0, "run", "all", "Wall-clock seconds for train + eval."),
    ("n_warnings",  0, "run", "all", "Config warnings (see config.json for the texts)."),
    ("error",       0, "run", "all", "Error message, if the run did not complete."),

    # ---------------- Key Metrics (section 2; order fixed by KEY_METRICS) ----------------
    ("acc_prior",   1, "accuracy", "all",
     "THE deployable accuracy: gates sampled from the prior p(z|x) (Eq. 6.2); the label is NOT used. Compare to chance and ceiling."),
    ("nll_raw",     1, "calibration", "all",
     "Negative log-likelihood before temperature scaling. Proper scoring rule: punishes wrong AND wrongly-confident. Lower is better."),
    ("ece_raw",     1, "calibration", "all",
     "Expected calibration error before temperature scaling. CAUTION: rewards under-confidence (dropout scores best and is useless)."),
    ("T_star",      1, "calibration", "all",
     "Temperature minimising validation NLL. ~1 = calibrated out of the box; <<1 = under-confident; >>1 = over-confident."),
    ("nll_cal",     1, "calibration", "all",
     "NLL after temperature scaling. The soundest single quality number next to acc_prior."),
    ("ece_cal",     1, "calibration", "all",
     "ECE after temperature scaling. CAUTION: still rewards under-confidence -- never read alone."),

    # ---------------- Other metrics ----------------
    ("brier_raw",   2, "calibration", "all",
     "Multiclass Brier score before temperature scaling: mean squared distance between the "
     "predicted probability vector and the one-hot truth, bounded in [0, 2]. A proper "
     "scoring rule like NLL, but FINITE when a confident prediction is wrong -- on the OOD "
     "ladder NLL is dominated by a handful of such cases. It is also the only one of the "
     "three that charges for probability mass on the WRONG classes: ECE bins only the "
     "largest probability, NLL reads only the true class."),
    ("brier_cal",   2, "calibration", "all",
     "Brier score after temperature scaling, counterpart of nll_cal and ece_cal."),
    ("acc_q_map",   2, "diagnostic", "all",
     "Label-conditioned: q is fed [x, onehot(y)] and picks its MAP config. Measures routing CAPACITY, not prediction. CAN EXCEED THE CEILING -- never compare it to one."),
    ("chance",      2, "reference", "all",
     "Majority-class accuracy = the lower bound. acc_prior must clearly beat this."),
    ("bayes_ceiling", 2, "reference", "blob",
     "Irreducible Bayes limit set by class overlap = the upper bound. acc_prior near it means near-perfect."),

    # -- real-data tasks (no h_epi column: that quantity is `bald`)
    ("acc_w0",      2, "run", "fashion,cifar10_resnet",
     "What the FROZEN base map W0(alpha) scores on its own, without any adapter. The real x axis of the alpha sweep: alpha is a unitless mixing weight, this is the competence it buys. ~0.90 at alpha=0 (the trained Out layer), ~chance at alpha=1."),
    # corrupted counterparts of the key metrics; T* is fitted in-distribution and carried over
    ("acc_ood",     2, "ood", "fashion,cifar10_resnet",
     "Accuracy on the OOD probe selected by fashion_probe at fashion_ood_fraction. Read TOGETHER with ece_raw_ood: a method can win ECE by being diffuse, and only a matching accuracy tells the two apart. Measured: the gated mixture's calibration gain arrives at UNCHANGED accuracy, MC dropout's does not."),
    ("ece_raw_ood", 2, "ood", "fashion,cifar10_resnet",
     "ECE on the selected OOD probe WITHOUT temperature scaling -- the counterpart of ece_raw, and the column comparable to Liu et al.'s Table 3 (their DNN 0.3435, software BNN 0.0918, spintronic BNN 0.1066 at the 50 % letter blend)."),
    ("brier_raw_ood", 2, "ood", "fashion,cifar10_resnet",
     "Brier score on the selected OOD probe, counterpart of brier_raw. Bounded, so unlike nll_raw_ood it cannot be dominated by a handful of confidently wrong images -- which is precisely what happens at the far end of the blend ladder."),
    ("nll_raw_ood", 2, "ood", "fashion,cifar10_resnet",
     "NLL on the selected OOD probe, counterpart of nll_raw. The proper scoring rule, and the metric to trust when ece_raw_ood and acc_ood disagree: a merely diffuse model cannot win it."),
    ("ece_cal_ood", 2, "ood", "fashion,cifar10_resnet",
     "ECE on the OOD probe after applying the IN-DISTRIBUTION T*. Measured: often WORSE than ece_raw_ood for the gated mixture, because its T* < 1 sharpens -- right in-distribution, wrong under shift. Post-hoc temperature scaling is the wrong tool for a model whose raw uncertainty is already calibrated."),
    ("nll_cal_ood", 2, "ood", "fashion,cifar10_resnet",
     "NLL on the OOD probe after the in-distribution T*. Same caveat as ece_cal_ood."),
    ("ece_ood_noise", 2, "ood", "fashion,cifar10_resnet",
     "ECE on the NOISE probe at the same fraction (additive geometry, unstructured contaminant). Tests whether the effect needs the contaminant to be image-like."),
    ("ece_ood_rotate", 2, "ood", "fashion,cifar10_resnet",
     "ECE on the ROTATE probe at the same fraction (nothing is superposed). Tests whether the effect needs additivity at all. Measured: the advantage over the deterministic baseline is LARGEST here. The three probes do not agree on a ranking, which is why a single-probe benchmark would be an artefact."),
    ("h_total",     2, "uncertainty", "fashion,cifar10_resnet",
     "Mean predictive entropy H(E[p]) in nats on the CLEAN test set (Liu et al. Eq. 9). Its epistemic part is `bald`, its aleatoric part h_alea."),
    ("h_alea",      2, "uncertainty", "fashion,cifar10_resnet",
     "Mean aleatoric uncertainty E[H(p)] in nats on the clean test set. h_total - h_alea = bald, so a large h_total with a small bald means the model is unsure but its sub-models AGREE about being unsure."),
    ("h_total_ood", 2, "uncertainty", "fashion,cifar10_resnet", "Predictive entropy on the OOD probe."),
    ("h_alea_ood",  2, "uncertainty", "fashion,cifar10_resnet", "Aleatoric uncertainty on the OOD probe."),
    ("bald_ood",    2, "uncertainty", "fashion,cifar10_resnet",
     "Epistemic uncertainty on the OOD probe -- the same quantity as `bald`, measured on the corrupted set. It should RISE with the corruption; that it does is the qualitative signature of Liu et al.'s Figure 5E."),
    ("n_trained",   2, "run", "all",
     "Trainable parameters of the recognition + adapter + field + coupling block (W0 excluded, it is frozen; J counted as its K(K-1)/2 free entries). K*r*D + K*C*r + n_pre*(D+C+1) + H*(D+1) + K*(H+1) + K(K-1)/2, with n_pre = K for mf and 2K-1 for tree."),
    ("n_frozen",    2, "run", "all",
     "Frozen parameters carried by the deployed model: W0 (C*D), plus any feature front-end (the LeNet trunk contributes 60 856). Read with n_trained when comparing against methods that train everything."),
    ("sel_at_0p5",  2, "calibration", "all",
     "Selective accuracy at 50% coverage: drop the most uncertain half, how good is the rest? Rising = the uncertainty is USEFUL (a sound metric)."),
    ("sel_at_0p75", 2, "calibration", "all", "Selective accuracy at 75% coverage."),
    ("sel_at_0p25", 2, "calibration", "all", "Selective accuracy at 25% coverage."),
    ("cal_acc",     2, "calibration", "all",
     "Accuracy on the calibration held-out set (a different set than acc_prior -- do not confuse the two)."),
    ("p_alloff",    2, "gate-prior", "all",
     "COLLAPSE ALARM: prior mass on z=0 (all gates off) -> W(z)=W0, prediction falls back to the frozen base. High = the run is meaningless."),
    ("p_onehot",    2, "gate-prior", "all",
     "Prior mass on exactly one active gate = crisp routing (vs a diffuse multi-on code)."),
    ("p_multi",     2, "gate-prior", "all",
     "Prior mass on >=2 active gates: a diffuse / redundant code (inflates BALD everywhere)."),
    ("gate_rate",   2, "gate-prior", "all",
     "Mean prior gate activation kappa; also what the dropout baseline is matched to for a fair comparison."),
    ("bald",        2, "uncertainty", "all",
     "Epistemic signal: disagreement among prior-sampled sub-models, H[mean]-mean[H]. CAUTION: magnitude rewards diffuseness -- it is NOT a quality score."),
    ("j_offdiag",   2, "mechanism", "all",
     "Mean off-diagonal coupling J. NEGATIVE = mutual inhibition ('one at a time') = the mechanism engaged. NOT a collapse indicator."),
    ("J_off_*",     2, "mechanism", "all", "Individual off-diagonal coupling J_kj."),
    ("acc_dropout", 2, "baseline", "all",
     "MC-dropout baseline (input-agnostic gates, kappa matched to the model's gate rate). The GAP to acc_prior is the 'gates buy accuracy' claim."),
    ("bald_dropout", 2, "baseline", "all",
     "BALD of the dropout baseline. Usually HIGHER than the model's -- diffuse gates disagree more. Not a quality score."),
    ("acc_allon",   2, "baseline", "routing,fashion,cifar10_resnet",
     "AllOn / merged-adapter baseline (no gating). Capped at the best fixed linear map. On the real-data tasks this is the LINEAR CEILING -- real data has no closed-form Bayes limit, so acc_prior is read against this instead of bayes_ceiling."),
    ("routing_diag_minus_off", 2, "mechanism", "routing,fashion,cifar10_resnet",
     "Routing quality: mean activation of the cells that SHOULD be on minus mean of those that should not. High = the right expert fires for the right rule. On the real-data tasks the index is the true CLASS, so it reads as 'expert k fires for class k'; only defined when K equals the number of indices."),
    ("routing_diag", 2, "mechanism", "routing,fashion,cifar10_resnet", "Mean activation of the routing-matrix cells that SHOULD be on (the diagonal E[z_k|rule k], or E[z_k|class k] on the real-data tasks)."),
    ("routing_off", 2, "mechanism", "routing,fashion,cifar10_resnet", "Mean activation of the cells that should be OFF."),
    ("acc_ctx",     2, "run", "routing", "Context strength CTX used (large = clean routing, small = ambiguous)."),
    ("tree_parents", 2, "recognition", "routing,fashion,cifar10_resnet",
     "The FINAL recognition-tree topology as a parent list (parents[k] = parent gate of k, -1 = root), read from the trained spec. tree/mf share the fixed chain -1,0,1,...; tree_max_span shows the LEARNED Chow-Liu tree, so a difference here proves the refit changed the topology. CAUTION (measured 2026-08-07): a difference is NOT evidence that real structure was found. The Chow-Liu target is built from trained parameters, which carry an echo of the warm-up topology -- training the same model under three different chain wirings returns each wiring back (100%/100%/89%). Judge structure by w_*_offdiag against a known scale and by resampling stability, not by this column. Text, not a number."),
    ("n_refits", 2, "recognition", "routing,fashion,cifar10_resnet",
     "How many Chow-Liu refits actually executed. 0 for tree/mf (fixed topology), >=1 for tree_max_span ('once' -> 1, 'periodic' -> more). The DIRECT proof that structure learning ran at all -- unlike tree_parents it stays informative when the learned tree happens to equal the warm-up chain."),
    ("tree_is_chain", 2, "recognition", "routing,fashion,cifar10_resnet",
     "1.0 if the final topology is the natural chain -1,0,1,...,K-2, else 0.0. On routing there is no ground-truth grouping to score a tree against (one expert per rule), so this replaces edge_recovery: it answers the only well-posed structural question here -- did the refit move the topology away from the warm-up chain? A 1.0 with n_refits>=1 means the MST re-derived the chain, NOT that the refit failed."),
    ("w_mean_offdiag", 2, "recognition", "routing,fashion,cifar10_resnet",
     "tree_max_span: mean off-diagonal Chow-Liu edge weight at the last refit -- the scale reference for w_block_contrast. On routing (no groups) it stands alone: it is how much pairwise posterior correlation the MST had to work with at all. Near 0 => the tree was fitted to noise, and whatever topology came out is arbitrary."),
    ("w_max_offdiag", 2, "recognition", "routing,fashion,cifar10_resnet",
     "tree_max_span: max off-diagonal Chow-Liu edge weight at the last refit -- the scale reference for w_block_contrast, and on routing the strength of the single most correlated gate pair."),
    ("acc_class_*", 2, "accuracy", "blob",
     "Prior-sampled accuracy for this true class. Reveals what the aggregate hides (R4: aggregate 0.76 while class 0 sat at 0.287)."),
]

_RUN_CSV_HEADER = ["metric", "value", "group", "description"]
SECTIONS = ("Config", "Key Metrics", "Other Metrics")   # matches tier 0 / 1 / 2 above

SECTION_FILL = {
    "Config": "D9E1F2", "Key Metrics": "FFE699",
    "accuracy": "E2EFDA", "reference": "F2F2F2", "calibration": "FCE4D6",
    "gate-prior": "DDEBF7", "mechanism": "D6DCE4", "uncertainty": "EAD1DC",
    "recognition": "E2F0D9", "redundancy": "F4CCCC", "baseline": "FFF2CC",
    "diagnostic": "D9D2E9", "run": "F2F2F2",
}


# ----------------------------------------------------------------- formatting
def sig(v, n=6):
    """Round to n significant digits (0.8916666666666667 -> 0.891667)."""
    if not isinstance(v, float) or v == 0 or not math.isfinite(v):
        return v
    return round(v, -int(math.floor(math.log10(abs(v)))) + (n - 1))


def fmt(v, decimal=","):
    """Format one value for the CSV (6 significant digits; locale decimal mark)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        s = f"{v:.6g}"
        return s.replace(".", decimal) if decimal != "." else s
    return str(v)


def _sep(locale):
    return (";", ",") if locale == "de" else (",", ".")


# ----------------------------------------------------------------- sectioning
def _nat(k):
    """Natural sort key: J_off_2 before J_off_10 (plain sort would invert them)."""
    m = re.search(r"(\d+)$", k)
    return (k[:m.start()], int(m.group(1))) if m else (k, -1)


def _expand(key, flat, used):
    if key.endswith("*"):
        return sorted((k for k in flat if k.startswith(key[:-1]) and k not in used), key=_nat)
    return [key] if key in flat else []


def blocks_for(flat, task):
    """[(section, [(metric, value, group, description), ...]), ...] in output order.
       Unregistered keys are dropped here."""
    cfg, key, other, used = [], [], [], set()
    for k, tier, group, applies, desc in ORDER:
        if applies != "all" and task not in applies.split(","):
            continue
        for kk in _expand(k, flat, used):
            used.add(kk)
            row = (kk, flat[kk], group, desc)
            (cfg if tier == 0 else key if kk in KEY_METRICS else other).append(row)
    key.sort(key=lambda r: KEY_METRICS.index(r[0]))   # headline metrics keep KEY_METRICS order
    return list(zip(SECTIONS, (cfg, key, other)))


def section_of(metric_key):
    """Which section a column belongs to (for the results.xlsx band row)."""
    if metric_key in KEY_METRICS:
        return "Key Metrics"
    for k, tier, group, _a, _d in ORDER:
        if k == metric_key or (k.endswith("*") and metric_key.startswith(k[:-1])):
            return "Config" if tier == 0 else group
    return "Other Metrics"


def column_order(records):
    """Config keys, then Key Metrics, then the rest in registry order, then extras.
       Union over all records; unregistered keys are appended, not dropped."""
    present = {k for r in records for k in r} - EXCLUDE
    known, used = [], set()
    for k, tier, _g, _a, _d in ORDER:
        if tier != 0:
            continue
        if k in present:
            known.append(k); used.add(k)
    for k in KEY_METRICS:
        if k in present:
            known.append(k); used.add(k)
    for k, tier, _g, _a, _d in ORDER:
        if tier == 0:
            continue
        for kk in (sorted((x for x in present if x.startswith(k[:-1])), key=_nat)
                   if k.endswith("*") else ([k] if k in present else [])):
            if kk not in used:
                known.append(kk); used.add(kk)
    return known + sorted(present - used)


# ----------------------------------------------------------------- per-run CSV
def write_run_csv(path, flat, task, locale="de"):
    """Long/tidy per-run CSV; a blank row between the sections."""
    delim, dec = _sep(locale)
    # utf-8-sig: the BOM makes Excel decode umlauts
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delim)
        w.writerow(_RUN_CSV_HEADER)
        for i, (_section, rows) in enumerate(blocks_for(flat, task)):
            if not rows:
                continue
            if i > 0:
                w.writerow([])                       # section separator
            for k, v, g, d in rows:
                w.writerow([k, fmt(v, dec), g, d])
    return path


# ----------------------------------------------------------------- results.xlsx
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _read_existing(path):
    """Existing rows of results.xlsx (band row 1, header row 2, data row 3+)."""
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path)
    ws = wb["Results"] if "Results" in wb.sheetnames else wb.active
    head = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(3, ws.max_row + 1):
        row = {head[c]: ws.cell(r, c + 1).value
               for c in range(len(head)) if head[c] is not None}
        if any(v is not None and str(v) != "" for v in row.values()):
            out.append(row)
    return out


def write_results_xlsx(path, records):
    """Wide, formatted, accumulating results workbook.  Upsert by run_id: existing rows
       of an incoming run_id are dropped first (all seeds), other runs are kept."""
    incoming_ids = {r.get("run_id") for r in records}
    kept = [r for r in _read_existing(path) if r.get("run_id") not in incoming_ids]
    allrec = kept + [{k: sig(v) if isinstance(v, float) else v for k, v in r.items()}
                     for r in records]
    cols = column_order(allrec)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    # band row: one merged, coloured header per run of same-section columns
    j = 1
    while j <= len(cols):
        sec = section_of(cols[j - 1])
        j2 = j
        while j2 < len(cols) and section_of(cols[j2]) == sec:
            j2 += 1
        ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=j2)
        c = ws.cell(1, j, sec)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=SECTION_FILL.get(sec, "FFFFFF"))
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER
        j = j2 + 1

    for i, k in enumerate(cols, start=1):
        h = ws.cell(2, i, k)
        h.font = Font(bold=True, size=9)
        h.fill = PatternFill("solid", fgColor="F2F2F2")
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = _BORDER
        ws.column_dimensions[get_column_letter(i)].width = (
            22 if k in ("run_id", "timestamp", "error") else max(9, min(16, len(k) + 3)))

    for r, rec in enumerate(allrec, start=3):
        for i, k in enumerate(cols, start=1):
            v = rec.get(k)
            cell = ws.cell(r, i, v if v != "" else None)
            cell.border = _BORDER
            if isinstance(v, float):
                cell.number_format = "0.000000"
            if k in KEY_METRICS:
                cell.font = Font(bold=True, size=10)

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(3, len(allrec) + 2)}"
    wb.save(path)
    return path, len(allrec), len(kept)
